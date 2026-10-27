"""LabDash laptop sync.

Downloads your encrypted LabDash from its private GitHub repository, decrypts it on this laptop and writes:
  * OneDrive/LabDash/LabDash-data.xlsx   (projects, tasks, log, samples, one sheet per protocol)
  * photos into each project's OneDrive folder, named by sample ID and date

Your passphrase and GitHub token are kept in Windows Credential Manager (via the keyring package),
never in a file. Run once with "setup", then a scheduled task runs "run" every 15 minutes.

    python labdash_sync.py setup     # asks for repo, token, passphrase, OneDrive folder; adds the scheduled task
    python labdash_sync.py run       # one sync (what the scheduled task does)
    python labdash_sync.py remove    # removes the scheduled task and stored secrets
"""
import base64, datetime as dt, getpass, json, os, re, subprocess, sys, tempfile, time, urllib.request, urllib.error
from pathlib import Path

APP = "LabDash"
STATE_DIR = Path(os.environ.get("LOCALAPPDATA", Path.home() / ".local" / "share")) / "LabDash"
STATE = STATE_DIR / "state.json"
LOG = STATE_DIR / "sync.log"


def log(msg):
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    line = dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S ") + msg
    print(line)
    with open(LOG, "a", encoding="utf-8") as f:
        f.write(line + "\n")
    try:  # keep the log small
        if LOG.stat().st_size > 400_000:
            LOG.write_text("\n".join(LOG.read_text(encoding="utf-8").splitlines()[-1500:]) + "\n", encoding="utf-8")
    except OSError:
        pass


def load_state():
    try:
        return json.loads(STATE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def save_state(s):
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    STATE.write_text(json.dumps(s, indent=1), encoding="utf-8")


# ---------- secrets ----------
def secret(name):
    import keyring
    return keyring.get_password(APP, name)


def set_secret(name, value):
    import keyring
    keyring.set_password(APP, name, value)


# ---------- GitHub ----------
class Source:
    def __init__(self, repo, token, local_dir=None):
        self.repo, self.token, self.local = repo, token, Path(local_dir) if local_dir else None

    def get(self, path):
        """Return the file's bytes, or None if it doesn't exist."""
        if self.local:
            p = self.local / path
            return p.read_bytes() if p.exists() else None
        req = urllib.request.Request(
            f"https://api.github.com/repos/{self.repo}/contents/{path}",
            headers={"Authorization": "Bearer " + self.token, "Accept": "application/vnd.github.raw+json",
                     "X-GitHub-Api-Version": "2022-11-28", "User-Agent": "LabDash-laptop-sync"})
        try:
            with urllib.request.urlopen(req, timeout=60) as r:
                return r.read()
        except urllib.error.HTTPError as e:
            if e.code == 404:
                return None
            if e.code == 401:
                raise SystemExit("GitHub rejected the token. It may have expired: run setup again with a new one.")
            raise

    def repo_info(self):
        if self.local:
            return {"private": True}
        req = urllib.request.Request(f"https://api.github.com/repos/{self.repo}",
                                     headers={"Authorization": "Bearer " + self.token, "Accept": "application/vnd.github+json",
                                              "User-Agent": "LabDash-laptop-sync"})
        with urllib.request.urlopen(req, timeout=30) as r:
            return json.loads(r.read())


# ---------- crypto (matches the app: PBKDF2-SHA256 -> AES-GCM-256 key wrap, AES-GCM data) ----------
def unlock(vault, passphrase):
    from cryptography.hazmat.primitives import hashes
    from cryptography.hazmat.primitives.kdf.pbkdf2 import PBKDF2HMAC
    from cryptography.hazmat.primitives.ciphers.aead import AESGCM
    from cryptography.exceptions import InvalidTag
    p = vault["keys"]["pass"]
    kdf = PBKDF2HMAC(algorithm=hashes.SHA256(), length=32, salt=base64.b64decode(p["salt"]), iterations=int(p["iter"]))
    wk = kdf.derive(passphrase.encode("utf-8"))
    try:
        raw = AESGCM(wk).decrypt(base64.b64decode(p["iv"]), base64.b64decode(p["ct"]), None)
    except InvalidTag:
        raise SystemExit("The stored passphrase doesn't open this LabDash. Run setup again.")
    key = AESGCM(raw)
    docs = json.loads(key.decrypt(base64.b64decode(vault["iv"]), base64.b64decode(vault["ct"]), None))
    return key, docs


def decrypt_image(key, rec):
    return key.decrypt(base64.b64decode(rec["iv"]), base64.b64decode(rec["ct"]), None)


# ---------- output ----------
def when(ms):
    return dt.datetime.fromtimestamp(ms / 1000) if isinstance(ms, (int, float)) and ms else None


def day(ms):
    w = when(ms)
    return w.date() if w else None


def day(iso):
    try:
        return dt.date.fromisoformat(iso) if iso else None
    except ValueError:
        return None


def safe(name):
    return re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(name)).strip(" .") or "item"


def project_dir(onedrive_root, p):
    folder = (p or {}).get("folder") or ""
    if folder.strip("/ "):
        parts = [safe(x) for x in folder.replace("\\", "/").split("/") if x.strip()]
        return onedrive_root.joinpath(*parts)
    return onedrive_root / "LabDash" / safe((p or {}).get("name") or "No project")


def derived(fields, values):
    def by(label):
        for f in fields:
            if f["label"].strip().lower() == label:
                try:
                    return float(str(values.get(f["key"], "")).replace(",", "."))
                except ValueError:
                    return None
        return None
    import math
    m, d, h, F = by("mass"), by("diameter"), by("height"), by("breaking force")
    dens = round(m / (math.pi * (d / 20) ** 2 * (h / 10)), 3) if m and d and h else None
    stren = round(1000 * F / (math.pi * (d / 2) ** 2), 1) if F and d else None
    return dens, stren


def write_workbook(docs, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    projects = {p["id"]: p for p in docs.get("projects", []) if not p.get("deleted")}
    protos = {p["id"]: p for p in docs.get("protocols", []) if not p.get("deleted")}
    pname = lambda i: (projects.get(i) or {}).get("name", "")
    wb = Workbook()
    head_font, head_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="166534")

    def sheet(title, header, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = safe(title)[:31]
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, dt.datetime):
                    c.number_format = "dd/mm/yyyy hh:mm"
                elif isinstance(c.value, dt.date):
                    c.number_format = "dd/mm/yyyy"
        ws.freeze_panes = "A2"
        for i, h in enumerate(header, 1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if i - 1 < len(r) and r[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width + 2))
        ws.auto_filter.ref = ws.dimensions
        return ws

    sheet("Projects", ["Project", "Summary", "Status", "Progress %", "Next step", "Folder", "Keywords", "Prefix"],
          [[p["name"], p.get("summary", ""), p.get("status", ""), p.get("progress", 0), p.get("nextStep", ""), p.get("folder", ""),
            p.get("keywords", ""), p.get("prefix", "")] for p in projects.values()], first=True)
    allp = {p["id"]: p for p in docs.get("projects", [])}
    pname_any = lambda i: (allp.get(i) or {}).get("name", "")

    def tstatus(t):
        if t.get("deleted"):
            return "Deleted"
        if t.get("done"):
            return "Done"
        if (allp.get(t.get("projectId")) or {}).get("status") == "Archived":
            return "Open (project archived)"
        if t.get("due") and t["due"] < dt.date.today().isoformat():
            return "Overdue"
        return "Open"
    order = {"Overdue": 0, "Open": 1, "Open (project archived)": 2, "Done": 3, "Deleted": 4}
    tasks = sorted(docs.get("tasks", []), key=lambda t: (order[tstatus(t)], t.get("due") or "9999"))
    sheet("Tasks", ["Task", "Project", "Status", "Due", "Completed", "Priority", "Created", "Last change", "Notes"],
          [[t["title"], pname_any(t.get("projectId")), tstatus(t), day(t.get("due")), day(t.get("completed")), t.get("priority", ""), when(t.get("created")),
            when(max([h.get("at", 0) for h in t.get("history") or []] or [0])), t.get("notes", "")] for t in tasks])
    events = sorted([(h.get("at", 0), t, h) for t in docs.get("tasks", []) for h in t.get("history") or []], key=lambda x: -x[0])
    sheet("Task history", ["When", "Task", "Project", "Change"], [[when(a), t["title"], pname_any(t.get("projectId")), h.get("what", "")] for a, t, h in events])
    sheet("Log", ["Logged", "Project", "Title", "Note", "Sample", "Photos"],
          [[when(l.get("when")), pname(l.get("projectId")), l.get("title", ""), l.get("note", ""), l.get("sampleId", ""), len(l.get("photos") or [])]
           for l in sorted(docs.get("log", []), key=lambda l: -(l.get("when") or 0)) if not l.get("deleted")])
    samples = sorted([s for s in docs.get("samples", []) if not s.get("deleted")], key=lambda s: s.get("created") or 0)
    sheet("Samples", ["SampleID", "Project", "Protocol", "Date", "Location", "Status", "Notes"],
          [[s["sampleId"], pname(s.get("projectId")), (protos.get(s.get("protocolId")) or {}).get("name", s.get("protocolName", "")), day(s.get("created")),
            s.get("location", ""), s.get("status", ""), s.get("notes", "")] for s in samples])
    for pr in protos.values():
        used = {k for s in samples if s.get("protocolId") == pr["id"] for k in (s.get("values") or {})}
        # fields removed in the app are left out unless a sample already holds a value for them
        fields = sorted([f for f in pr.get("fields", []) if not f.get("off") or f["key"] in used], key=lambda f: f.get("step", 0))
        rows = []
        for s in samples:
            if s.get("protocolId") != pr["id"]:
                continue
            v = s.get("values") or {}
            dens, stren = derived(fields, v)
            rows.append([s["sampleId"], day(s.get("created")), s.get("location", "")] + [v.get(f["key"]) for f in fields] + [dens, stren, s.get("notes", "")])
        header = ["SampleID", "Date", "Location"] + [f["label"] + (f" ({f['unit']})" if f.get("unit") else "") for f in fields] + ["Density (g/cm³)", "Strength F/A (kPa)", "Notes"]
        sheet(pr["name"], header, rows)
    sheet("Links", ["Project", "Name", "Link", "Added"],
          [[p["name"], k.get("title", ""), k.get("url", ""), when(k.get("added"))] for p in projects.values() for k in p.get("links") or []])
    for ws in wb.worksheets:
        if ws.title == "Links":
            for r in ws.iter_rows(min_row=2, min_col=3, max_col=3):
                for c in r:
                    if c.value:
                        c.hyperlink = c.value
                        c.style = "Hyperlink"
    sheet("Locations", ["Location"], [[l["name"]] for l in docs.get("locations", [])])
    path.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=path.parent)
    os.close(fd)
    wb.save(tmp)
    try:
        os.replace(tmp, path)
        return True
    except PermissionError:
        os.remove(tmp)
        log("LabDash-data.xlsx is open in Excel, so it wasn't updated this time. Close it and the next sync will update it.")
        return False


class RunLock:
    """Stops two syncs (the scheduled one and a manual one) running at once and saving photos twice."""
    def __init__(self):
        self.path = STATE_DIR / "sync.lock"
        self.held = False

    def __enter__(self):
        STATE_DIR.mkdir(parents=True, exist_ok=True)
        try:
            if self.path.exists() and time.time() - self.path.stat().st_mtime > 600:
                self.path.unlink()  # left behind by a sync that crashed
            fd = os.open(self.path, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
            os.write(fd, str(os.getpid()).encode())
            os.close(fd)
            self.held = True
        except FileExistsError:
            pass
        return self

    def __exit__(self, *a):
        if self.held:
            try:
                self.path.unlink()
            except OSError:
                pass


def sync(src, passphrase, onedrive_root):
    with RunLock() as lock:
        if not lock.held:
            log("Another sync is already running, so this one was skipped.")
            return
        _sync(src, passphrase, onedrive_root)


def _sync(src, passphrase, onedrive_root):
    raw = src.get("vault.json")
    if raw is None:
        log("No LabDash in the repository yet. Set it up on your phone first.")
        return
    key, docs = unlock(json.loads(raw), passphrase)
    st = load_state()
    done = set(st.get("photos", []))
    projects = {p["id"]: p for p in docs.get("projects", [])}
    new = 0
    for l in docs.get("log", []):
        if l.get("deleted"):
            continue
        stamp = (when(l.get("when")) or dt.datetime.now()).strftime("%Y-%m-%d_%H%M")
        base = safe(l.get("sampleId") or l.get("title") or "photo")[:60]
        for n, pid in enumerate(l.get("photos") or [], 1):
            if pid in done:
                continue
            blob = src.get(f"img/{pid}.json")
            if blob is None:
                continue  # not uploaded yet; try next time
            folder = project_dir(onedrive_root, projects.get(l.get("projectId")))
            folder.mkdir(parents=True, exist_ok=True)
            name = f"{base}_{stamp}" + (f"_{n}" if len(l["photos"]) > 1 else "") + ".jpg"
            data = decrypt_image(key, json.loads(blob))
            target = folder / name
            i = 2
            while target.exists() and target.read_bytes() != data:
                target = folder / f"{Path(name).stem}~{i}.jpg"
                i += 1
            if not target.exists():
                target.write_bytes(data)
                new += 1
            done.add(pid)
            st["photos"] = sorted(done)
            save_state(st)  # record each photo straight away
    wrote = write_workbook(docs, onedrive_root / "LabDash" / "LabDash-data.xlsx")
    st["photos"] = sorted(done)
    st["last"] = time.time()
    save_state(st)
    log(f"Synced: {len(docs.get('samples', []))} samples, {len(docs.get('tasks', []))} tasks, {new} new photos" + ("" if wrote else " (workbook skipped)"))


def default_onedrive():
    for k in ("OneDriveConsumer", "OneDrive", "OneDriveCommercial"):
        if os.environ.get(k):
            return os.environ[k]
    return str(Path.home() / "OneDrive")


def setup():
    print("LabDash laptop sync setup\n")
    repo = input("Private data repository (e.g. your-username/LabDash-data): ").strip()
    repo = re.sub(r"^https?://github\.com/", "", repo).rstrip("/").removesuffix(".git")
    print("Paste the GitHub token below. In this window, paste with a RIGHT-CLICK (Ctrl+V may not work), then press Enter.")
    token = input("GitHub access token: ").strip()
    if not token.startswith("github_pat_"):
        raise SystemExit("That doesn't look like a fine-grained token (they start with github_pat_). Run install.bat again and paste it with a right-click.")
    src = Source(repo, token)
    try:
        info = src.repo_info()
    except urllib.error.HTTPError as e:
        raise SystemExit(f"GitHub said {e.code}. Check the repository name and that the token can read it.")
    if not info.get("private"):
        raise SystemExit("That repository is public. Make it private on GitHub first.")
    passphrase = getpass.getpass("LabDash passphrase (type it; nothing shows as you type, then press Enter): ")
    raw = src.get("vault.json")
    if raw is None:
        raise SystemExit("There's no LabDash in that repository yet. Set it up on your phone and sync first.")
    unlock(json.loads(raw), passphrase)
    od = input(f"OneDrive folder [{default_onedrive()}]: ").strip() or default_onedrive()
    set_secret("repo", repo)
    set_secret("token", token)
    set_secret("passphrase", passphrase)
    set_secret("onedrive", od)
    print("\nSaved to Windows Credential Manager. Running a first sync...")
    sync(src, passphrase, Path(od))
    if os.name == "nt":
        pyw = Path(sys.executable).with_name("pythonw.exe")
        exe = pyw if pyw.exists() else Path(sys.executable)
        cmd = f'"{exe}" "{Path(__file__).resolve()}" run'
        r = subprocess.run(["schtasks", "/Create", "/F", "/SC", "MINUTE", "/MO", "15", "/TN", "LabDash sync", "/TR", cmd], capture_output=True, text=True)
        print("Scheduled task added: runs every 15 minutes." if r.returncode == 0 else "Couldn't add the scheduled task:\n" + r.stdout + r.stderr)
    print(f"\nDone. Your workbook: {Path(od) / 'LabDash' / 'LabDash-data.xlsx'}\nLog: {LOG}")


def run():
    repo, token, pw, od = secret("repo"), secret("token"), secret("passphrase"), secret("onedrive")
    if not (repo and token and pw and od):
        raise SystemExit("Not set up yet: run  python labdash_sync.py setup")
    try:
        sync(Source(repo, token), pw, Path(od))
    except urllib.error.URLError as e:
        log(f"Offline or GitHub unreachable ({e}). Will try again next time.")


def remove():
    import keyring
    for n in ("repo", "token", "passphrase", "onedrive"):
        try:
            keyring.delete_password(APP, n)
        except Exception:
            pass
    if os.name == "nt":
        subprocess.run(["schtasks", "/Delete", "/F", "/TN", "LabDash sync"])
    print("Removed the scheduled task and stored secrets. Your OneDrive files are untouched.")


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else "run"
    if cmd == "setup":
        try:
            setup()
        except KeyboardInterrupt:
            print("\nSetup was cancelled (that happens if Ctrl+C is pressed). Run install.bat again.")
        except SystemExit as e:
            if e.code and not isinstance(e.code, int):
                print("\nSetup stopped: " + str(e.code))
        except Exception as e:
            print(f"\nSetup hit an error: {e!r}\nSend Claude a screenshot of this window.")
    elif cmd == "run":
        try:
            run()
        except SystemExit as e:
            if e.code and not isinstance(e.code, int):
                log(str(e.code))
            raise
        except Exception as e:  # keep a record when running hidden from the scheduler
            log(f"Error: {e!r}")
            raise
    elif cmd == "remove":
        remove()
    elif cmd == "test-local":  # developer check: python labdash_sync.py test-local <repo-dir> <passphrase> <out-dir>
        sync(Source("", "", sys.argv[2]), sys.argv[3], Path(sys.argv[4]))
    else:
        print(__doc__)
